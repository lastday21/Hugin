from __future__ import annotations

import json
from collections.abc import Iterator
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy.orm import Session

from hugin.core.settings import Settings
from hugin.database import create_database
from hugin.database.models import CandidateProfileModel, ResumeModel
from hugin.repositories import AccountRepository, ResumeRepository
from hugin.services.resume_improvement import ResumeImprovementService


class LocalResumeModel:
    model_name = "local-resume-review"

    def __init__(self, outputs: tuple[str, ...], question: str | None = None) -> None:
        self.outputs = iter(outputs)
        self.question = question
        self.rewrite_prompts: list[str] = []

    def complete(self, _system_prompt: str, prompt: str) -> str:
        if "JSON-массив" in prompt:
            return json.dumps(
                [
                    {
                        "topic": topic,
                        "known": not (topic == "scale" and self.question),
                        "evidence": None
                        if topic == "scale" and self.question
                        else "Исходное описание",
                        "question": self.question if topic == "scale" else None,
                    }
                    for topic in (
                        "personal_contribution",
                        "result",
                        "project_status",
                        "scale",
                        "collaboration",
                    )
                ],
                ensure_ascii=False,
            )
        self.rewrite_prompts.append(prompt)
        return next(self.outputs)


def resume_source(narrative: str) -> str:
    return (
        "Опыт работы\nЯнварь 2024 — настоящее время\nКомпания\nРазработчик\n"
        f"{narrative}\nОбразование\nУниверситет\n"
    )


@pytest.fixture
def resume_record(settings: Settings) -> Iterator[tuple[Session, int, ResumeModel]]:
    database = create_database(settings)
    try:
        with database.sessions.begin() as session:
            account = AccountRepository(session).create("Проверка", "independent-resume")
            record = ResumeRepository(session).upsert(account.id, "review-resume", "Python")
            resume = session.get(ResumeModel, record.id)
            assert resume is not None
            session.add(
                CandidateProfileModel(
                    account_id=account.id,
                    active_resume_id=resume.id,
                    display_name="Проверка",
                )
            )
            session.flush()
            yield session, account.id, resume
    finally:
        database.close()


@pytest.mark.integration
@pytest.mark.parametrize(
    ("source", "rewrite", "question", "answer"),
    [
        pytest.param(
            "Проект Alpha:\n- Обработал 73 заявки за месяц.",
            "Проект Beta:\n- Обработал 73 заявки за месяц.",
            None,
            "",
            id="explicit-project-substitution",
        ),
        pytest.param(
            "Проект Alpha:\n- Разрабатывал сервис.",
            "Результаты:\n- Обработал 73 заявки за месяц.",
            "Сколько заявок обработал этот проект?",
            "В проекте Beta обработал 73 заявки за месяц.",
            id="foreign-project-answer-with-output-heading-omitted",
        ),
        pytest.param(
            "- На тестовом стенде:\n- Обработал 73 заявки за месяц.",
            "- Обработал 73 заявки за месяц.",
            None,
            "",
            id="testing-condition-in-previous-line",
        ),
        pytest.param(
            "- За январь:\n- Обработал 73 заявки.",
            "- За февраль обработал 73 заявки.",
            None,
            "",
            id="calendar-heading-change",
        ),
        pytest.param(
            "- За январь:\n- Обработал 73 заявки.",
            "- Обработал 73 заявки.",
            None,
            "",
            id="calendar-heading-omitted",
        ),
        pytest.param(
            "- Время обработки составляет 73 секунды.",
            "- Время обработки составляет 73 минуты.",
            None,
            "",
            id="seconds-are-not-minutes",
        ),
        pytest.param(
            "- Обработал не менее 73 заявок за месяц.",
            "- Обработал ровно 73 заявки за месяц.",
            None,
            "",
            id="lower-bound-is-not-exact-count",
        ),
        pytest.param(
            "- Обработал 73 заявки за неделю.",
            "- Обработал 73 заявки за месяц.",
            None,
            "",
            id="weekly-count-is-not-monthly",
        ),
        pytest.param(
            "- Разрабатывал сервис.",
            "- Обработал 73 заявки за месяц.",
            "Вы обработали 73 заявки за месяц?",
            "Таких измерений не проводил.",
            id="question-is-not-evidence",
        ),
        pytest.param(
            "- Разрабатывал сервис.",
            "- Обработал 73 заявки за месяц.",
            None,
            "",
            id="target-and-vacancies-are-not-evidence",
        ),
    ],
)
def test_unverified_numeric_fact_is_rejected_before_any_artifact(
    resume_record: tuple[Session, int, ResumeModel],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    source: str,
    rewrite: str,
    question: str | None,
    answer: str,
) -> None:
    session, account_id, resume = resume_record
    original = resume_source(source)
    resume.content_text = original
    model = LocalResumeModel((rewrite, rewrite), question)
    service = ResumeImprovementService(session, tmp_path, model)
    monkeypatch.setattr(
        service,
        "_vacancy_context",
        lambda *_: "Обработал 73 заявки за месяц — требование вакансий.",
    )
    with pytest.raises(ValueError, match="Число не подтверждено"):
        service.improve(
            account_id,
            lambda *_: answer,
            target_role="Разработчик: обработал 73 заявки за месяц",
        )
    assert resume.content_text == original
    assert len(model.rewrite_prompts) == 2
    assert not list(tmp_path.rglob("*.docx"))
    assert not list(tmp_path.rglob("*.json"))


@pytest.mark.integration
@pytest.mark.parametrize("project", [False, True])
@pytest.mark.parametrize("repair", [False, True])
def test_confirmed_answer_is_bound_to_current_block_and_can_repair_once(
    resume_record: tuple[Session, int, ResumeModel],
    tmp_path: Path,
    project: bool,
    repair: bool,
) -> None:
    session, account_id, resume = resume_record
    heading = "Проект Alpha:\n" if project else ""
    original = resume_source(f"{heading}- Разрабатывал сервис.")
    resume.content_text = original
    valid = f"{heading}- Обработал 73 заявки за месяц."
    invalid = f"{heading}- Обработал 730 заявок за месяц."
    model = LocalResumeModel(
        (invalid, valid) if repair else (valid, valid),
        "Сколько заявок обработал этот проект?",
    )
    result = ResumeImprovementService(session, tmp_path, model).improve(
        account_id, lambda *_: "Обработал 73 заявки за месяц."
    )
    paragraphs = "\n".join(
        paragraph.text for paragraph in Document(str(result.draft_path)).paragraphs
    )
    assert "73 заявки" in paragraphs
    assert "730" not in paragraphs
    assert resume.content_text == original
    assert result.source_unchanged
    assert len(model.rewrite_prompts) == (2 if repair else 1)
    report = json.loads(result.report_path.read_text(encoding="utf-8"))
    assert len(report["blocks"]) == 1
    assert report["blocks"][0]["answers"] == ["Обработал 73 заявки за месяц."]


@pytest.mark.integration
def test_testing_condition_can_move_from_heading_to_same_sentence(
    resume_record: tuple[Session, int, ResumeModel], tmp_path: Path
) -> None:
    session, account_id, resume = resume_record
    original = resume_source("- На тестовом стенде:\n- Обработал 73 заявки за месяц.")
    resume.content_text = original
    valid = "- На тестовом стенде обработал 73 заявки за месяц."
    model = LocalResumeModel((valid, valid))
    result = ResumeImprovementService(session, tmp_path, model).improve(account_id, lambda *_: "")
    assert len(model.rewrite_prompts) == 1
    assert resume.content_text == original
    rendered = "\n".join(p.text for p in Document(str(result.draft_path)).paragraphs)
    assert "На тестовом стенде обработал 73" in rendered


@pytest.mark.integration
def test_later_block_cannot_borrow_previous_project_count_or_leave_partial_draft(
    resume_record: tuple[Session, int, ResumeModel], tmp_path: Path
) -> None:
    session, account_id, resume = resume_record
    first = "Проект Alpha:\n- Обработал 73 заявки за месяц."
    second = "Проект Beta:\n- Разрабатывал сервис."
    invalid_second = "Проект Beta:\n- Обработал 73 заявки за месяц."
    original = resume_source(f"{first}\n{second}")
    resume.content_text = original
    model = LocalResumeModel((first, invalid_second, invalid_second))
    with pytest.raises(ValueError, match="Число не подтверждено"):
        ResumeImprovementService(session, tmp_path, model).improve(account_id, lambda *_: "")
    assert len(model.rewrite_prompts) == 3
    assert resume.content_text == original
    assert not list(tmp_path.rglob("*.docx"))
    assert not list(tmp_path.rglob("*.json"))
